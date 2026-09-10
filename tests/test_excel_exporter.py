from datetime import date

from openpyxl import load_workbook

from app.domain.models import PublicationStatus, SchedulePublication
from app.exporters import excel as excel_module
from app.exporters.excel import ExcelExporter


def test_publication_workbook_is_structured_and_formula_driven(tmp_path, monkeypatch):
    monkeypatch.setattr(excel_module, "EXPORTS_DIR", tmp_path)
    publication = SchedulePublication(
        version_number="EDT-2030-S02-V1",
        week_start=date(2030, 1, 7),
        week_end=date(2030, 1, 13),
        status=PublicationStatus.published,
        created_by="admin-test",
        snapshot_json=[
            {
                "entry_date": "2030-01-07",
                "timeslot_id": 1,
                "timeslot_label": "08:00 - 10:00",
                "group_id": 1,
                "group_name": "M2 RSI",
                "course_title": "Intelligence artificielle",
                "teacher_name": "Enseignant test",
                "room_name": "Salle A",
                "is_new": True,
            }
        ],
        generation_report={},
    )

    result = ExcelExporter().generate_publication_workbook(publication)
    workbook = load_workbook(result["path"], data_only=False)

    assert workbook.sheetnames == ["Synthèse", "Planning complet", "M2 RSI"]
    assert workbook["Synthèse"]["B10"].value.startswith("=COUNTA")
    assert workbook["Planning complet"]["A5"].is_date
    assert workbook["Planning complet"]["G5"].value == "Nouvelle séance"
    assert workbook["Planning complet"].freeze_panes == "A5"
