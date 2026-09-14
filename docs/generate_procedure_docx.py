"""Genere le document Word de la procedure de bascule UAT DC1 -> DC2."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).with_name("Procedure_bascule_UAT_DC1_DC2.docx")

NAVY = RGBColor(0x1F, 0x35, 0x64)
SLATE = RGBColor(0x44, 0x4C, 0x5A)
HEADER_FILL = "1F3564"
BAND_FILL = "EEF1F6"
NOTE_FILL = "FBF3E2"


# --------------------------------------------------------------------------- #
# Primitives de mise en forme
# --------------------------------------------------------------------------- #
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def setup_styles(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Heading 1", 15, NAVY, 18, 8),
        ("Heading 2", 12.5, NAVY, 14, 6),
        ("Heading 3", 11, SLATE, 10, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(3)


def setup_page(document):
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    return section


def build_footer(section, reference):
    footer = section.footer
    para = footer.paragraphs[0]
    para.text = ""
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(f"{reference}  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = SLATE
    add_field(para, "PAGE")
    para.add_run(" / ")
    add_field(para, "NUMPAGES")
    for run in para.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = SLATE


def build_header(section, title):
    para = section.header.paragraphs[0]
    para.text = title
    para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in para.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = SLATE
        run.font.name = "Calibri"


# --------------------------------------------------------------------------- #
# Blocs de contenu
# --------------------------------------------------------------------------- #
def para(document, text, style=None, bold=False, size=None, align=None,
         space_after=None):
    p = document.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(document, items):
    for item in items:
        p = document.add_paragraph(style="List Bullet")
        if isinstance(item, tuple):
            lead, rest = item
            run = p.add_run(lead)
            run.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def table(document, headers, rows, widths=None, band=True):
    tbl = document.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    head = tbl.rows[0]
    repeat_header(head)
    for idx, label in enumerate(headers):
        cell = head.cells[idx]
        cell.text = ""
        run = cell.paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].paragraph_format.space_after = Pt(2)
        cell.paragraphs[0].paragraph_format.space_before = Pt(2)
        shade(cell, HEADER_FILL)

    for r_idx, row_values in enumerate(rows):
        row = tbl.add_row()
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.space_before = Pt(2)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9.5)
            if band and r_idx % 2 == 1:
                shade(cell, BAND_FILL)

    if widths:
        for row in tbl.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    return tbl


def note_box(document, title, lines):
    tbl = document.add_table(rows=1, cols=1)
    tbl.style = "Table Grid"
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(16.6)
    shade(cell, NOTE_FILL)
    cell.text = ""

    head = cell.paragraphs[0]
    head.paragraph_format.space_after = Pt(4)
    run = head.add_run(title)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = NAVY

    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(line)
        run.font.size = Pt(9.5)
    document.add_paragraph()
    return tbl


def phase_block(document, number, title, owner, objective, actions, controls,
                exits):
    document.add_heading(f"{number} {title}", level=2)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("Responsable : ")
    run.bold = True
    run.font.size = Pt(9.5)
    run = p.add_run(owner)
    run.font.size = Pt(9.5)

    para(document, objective)

    document.add_heading("Actions", level=3)
    bullets(document, actions)

    document.add_heading("Contrôles", level=3)
    bullets(document, controls)

    document.add_heading("Critères de sortie", level=3)
    bullets(document, exits)


# --------------------------------------------------------------------------- #
# Document
# --------------------------------------------------------------------------- #
REFERENCE = "PROC-UAT-DC1-DC2 v1.0"
TITLE = "Procédure de bascule de l'environnement UAT"
SUBTITLE = ("Restauration DC1 vers DC2, mise en service, réplication "
            "et tests de failover / failback")


def cover(document):
    for _ in range(3):
        document.add_paragraph()

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("DOCUMENT TECHNIQUE DE PROCÉDURE")
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = SLATE

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(TITLE)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = NAVY

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run(SUBTITLE)
    run.font.size = Pt(12)
    run.font.color.rgb = SLATE

    table(
        document,
        ["Rubrique", "Valeur"],
        [
            ["Référence", "PROC-UAT-DC1-DC2"],
            ["Version", "1.0"],
            ["Date", "04/09/2026"],
            ["Objet", "Séquence de mise en place de l'UAT sur DC2 et "
                      "validation des mécanismes de bascule"],
            ["Environnement", "UAT — sites DC1 (source) et DC2 (cible)"],
            ["Statut", "Pour validation"],
            ["Destinataires", "Client, équipe Infrastructure, équipe Base de "
                              "données, équipe Applicative"],
            ["Classification", "Confidentiel — diffusion restreinte"],
        ],
        widths=[4.6, 12.0],
    )

    document.add_paragraph()
    document.add_heading("Historique des versions", level=2)
    table(
        document,
        ["Version", "Date", "Auteur", "Nature de la modification"],
        [["1.0", "04/09/2026", "Équipe Base de données", "Création"]],
        widths=[2.2, 2.6, 4.4, 7.4],
    )

    document.add_paragraph()
    document.add_heading("Sommaire", level=2)
    p = document.add_paragraph()
    add_field(p, r'TOC \o "1-2" \h \z \u')
    p = document.add_paragraph()
    run = p.add_run("Sommaire à actualiser dans Word : sélectionner le champ "
                    "puis appuyer sur F9.")
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = SLATE

    document.add_page_break()


def section_objet(document):
    document.add_heading("1. Objet et périmètre", level=1)
    para(document,
         "Ce document décrit la séquence à suivre pour construire "
         "l'environnement UAT du site DC2 à partir de l'UAT du site DC1, puis "
         "pour valider la capacité de bascule entre les deux sites. Il précise "
         "pour chaque phase les actions à mener, les contrôles à réaliser et "
         "les critères permettant d'engager la phase suivante.")

    document.add_heading("1.1 Périmètre couvert", level=2)
    bullets(document, [
        "Sauvegarde et restauration des machines virtuelles de base de "
        "données, de DC1 vers DC2, avec réaffectation du plan d'adressage.",
        "Mise en service et contrôle d'intégrité de la base de données DC2.",
        "Déploiement et recette des applications sur DC2.",
        "Activation et validation de la réplication de flux de DC1 vers DC2.",
        "Exécution des tests de failover et de failback.",
        "Passage des applications DC2 en mode inactif.",
    ])

    document.add_heading("1.2 Périmètre exclu", level=2)
    bullets(document, [
        "Dimensionnement et acquisition de l'infrastructure DC2.",
        "Mise en place d'une adresse virtuelle ou d'un répartiteur de "
        "connexions ; la bascule des accès applicatifs reste manuelle.",
        "Bascule automatique sans intervention humaine.",
        "Modification du code des applications.",
        "Tests de charge et de performance.",
    ])


def section_sequence(document):
    document.add_heading("2. Vue d'ensemble de la séquence", level=1)
    para(document,
         "Les six phases s'exécutent dans l'ordre. Un point de décision "
         "formel est positionné entre la phase 3 et la phase 4.")
    table(
        document,
        ["Phase", "Intitulé", "Responsable", "Livrable"],
        [
            ["1", "Sauvegarde et restauration des VM de base de données "
                  "DC1 vers DC2, avec remap IP",
             "Infrastructure", "VM restaurées et adressées sur DC2"],
            ["2", "Mise en service de la base de données UAT DC2",
             "Base de données", "Instance DC2 démarrée et contrôlée"],
            ["3", "Déploiement et validation des applications sur DC2",
             "Applicatif", "Recette applicative DC2 prononcée"],
            ["—", "Point de décision : arbitrage et gel avant reconstruction",
             "Client / Base de données / Applicatif",
             "Autorisation écrite de poursuivre"],
            ["4", "Activation et validation de la réplication DC1 vers DC2",
             "Base de données", "Réplication active et vérifiée"],
            ["5", "Tests de failover et de failback",
             "Base de données / Applicatif",
             "Procès-verbal de test, durées mesurées"],
            ["6", "Passage des applications DC2 en mode inactif",
             "Applicatif", "DC2 en attente, réplication maintenue"],
        ],
        widths=[1.4, 7.0, 4.0, 4.2],
    )

    document.add_heading("2.1 État de l'environnement au fil des phases",
                         level=2)
    table(
        document,
        ["Après la phase", "Base de données DC2", "Applications DC2",
         "Lien DC1 – DC2"],
        [
            ["1", "Arrêtée", "Non déployées", "Aucun"],
            ["2", "Active, autonome, en lecture et écriture", "Non déployées",
             "Aucun"],
            ["3", "Active, autonome, en lecture et écriture",
             "Déployées et en service", "Aucun"],
            ["4", "Active, réplica en lecture seule",
             "Arrêtées ou raccordées à DC1", "Réplication DC1 vers DC2"],
            ["5", "Rôle inversé pendant le test, puis rétabli",
             "Raccordées au site actif", "Réplication rétablie après retour"],
            ["6", "Active, réplica en lecture seule", "Inactives",
             "Réplication DC1 vers DC2"],
        ],
        widths=[2.6, 5.2, 4.4, 4.4],
    )


def section_roles(document):
    document.add_heading("3. Rôles et responsabilités", level=1)
    para(document,
         "R : responsable de l'exécution — C : consulté — I : informé. "
         "Chaque passage de phase est tracé par un accord écrit.")
    table(
        document,
        ["Phase", "Infrastructure", "Base de données", "Applicatif", "Client"],
        [
            ["1. Restauration des VM et remap IP", "R", "C", "I", "I"],
            ["2. Mise en service de la base DC2", "C", "R", "C", "I"],
            ["3. Déploiement et recette applicative", "C", "C", "R", "C"],
            ["Point de décision", "C", "R", "R", "R"],
            ["4. Réplication DC1 vers DC2", "I", "R", "C", "I"],
            ["5. Failover et failback", "C", "R", "R", "C"],
            ["6. Applications inactives", "I", "C", "R", "I"],
        ],
        widths=[7.0, 2.5, 2.6, 2.3, 2.2],
    )


def section_prerequis(document):
    document.add_heading("4. Prérequis", level=1)
    para(document,
         "Les éléments suivants sont réunis et validés avant le lancement de "
         "la phase 1.")
    bullets(document, [
        ("Fenêtre d'intervention. ",
         "Créneau validé pour chaque phase, avec les interlocuteurs des trois "
         "équipes disponibles."),
        ("Sauvegarde source. ",
         "Sauvegarde récente et restaurable de l'UAT DC1, indépendante de la "
         "présente intervention."),
        ("Plan d'adressage DC2. ",
         "Adresses, masque, passerelle, résolution de noms et règles de "
         "filtrage réseau entre les deux sites arrêtés par écrit."),
        ("Capacité. ",
         "Espace disque, mémoire et processeurs de DC2 au moins équivalents à "
         "ceux de DC1, avec la marge nécessaire aux journaux de "
         "transactions."),
        ("Accès et comptes. ",
         "Comptes d'administration système et base de données, accès distant "
         "entre les deux serveurs, comptes de service applicatifs."),
        ("Inventaire des raccordements. ",
         "Liste exhaustive des chaînes de connexion des applications, "
         "condition indispensable à la bascule des accès en phase 5."),
        ("Jeu de recette. ",
         "Scénarios de test applicatifs et critères d'acceptation définis par "
         "le client, rejouables à l'identique."),
        ("Sauvegarde de la cible. ",
         "Si l'environnement DC2 existant contient des données à conserver, "
         "leur sauvegarde est réalisée avant la phase 1."),
    ])


def section_phases(document):
    document.add_heading("5. Déroulement détaillé", level=1)

    phase_block(
        document,
        "5.1", "Phase 1 — Sauvegarde et restauration des VM de base de "
               "données, avec remap IP",
        "Infrastructure, avec l'outil de sauvegarde et de restauration "
        "d'images de machines virtuelles",
        "Obtenir sur DC2 une copie fidèle des machines virtuelles de base de "
        "données de DC1, raccordée au réseau du site DC2.",
        actions=[
            "Contrôler, avant sauvegarde, l'espace disponible et l'état de "
            "fonctionnement de l'instance source.",
            "Exécuter la sauvegarde des machines virtuelles de base de "
            "données de DC1 et consigner l'identifiant du travail ainsi que "
            "l'horodatage du point de restauration retenu.",
            "Restaurer les machines sur l'infrastructure DC2 en réseau isolé, "
            "ou interfaces déconnectées, afin d'éviter tout conflit "
            "d'adresses avec les machines de DC1.",
            "Neutraliser le démarrage automatique du service de base de "
            "données : le premier démarrage est encadré par la phase 2.",
            "Réaffecter le plan d'adressage DC2 : adresse, masque, "
            "passerelle, résolution de noms et, le cas échéant, nom d'hôte.",
            "Recenser les configurations recopiées depuis DC1 qui référencent "
            "encore le site source, et les transmettre à l'équipe base de "
            "données.",
        ],
        controls=[
            "Adresse cible active et jointe depuis le réseau d'administration.",
            "Absence de conflit d'adresses avec les machines de DC1.",
            "Résolution de noms et synchronisation horaire correctes.",
            "Service de base de données arrêté, démarrage automatique "
            "désactivé.",
        ],
        exits=[
            "Machines virtuelles restaurées sur DC2 et adressées.",
            "Point de restauration utilisé consigné dans le journal "
            "d'exécution.",
            "Aucune interaction possible entre la copie et l'environnement "
            "DC1.",
        ],
    )

    phase_block(
        document,
        "5.2", "Phase 2 — Mise en service de la base de données UAT DC2",
        "Équipe Base de données",
        "Démarrer l'instance restaurée et démontrer que les bases de données "
        "sont exploitables. À l'issue de cette phase, DC2 constitue une "
        "instance autonome, en lecture et écriture, sans lien avec DC1.",
        actions=[
            "Avant le premier démarrage, neutraliser tout mécanisme hérité de "
            "DC1 susceptible d'écrire vers une destination partagée : "
            "archivage des journaux de transactions, tâches planifiées, "
            "agents de sauvegarde, outillage de gestion de cluster.",
            "Vérifier que l'instance est configurée pour démarrer en mode "
            "normal et non en mode restauration.",
            "Démarrer le service de base de données et analyser le journal de "
            "démarrage. Un rejeu des journaux de transactions est attendu "
            "lorsque la sauvegarde a été réalisée à chaud.",
            "Contrôler les paramètres effectifs de l'instance : port d'écoute, "
            "adresse d'écoute, emplacement des données.",
            "Établir l'inventaire des bases de données et relever leur "
            "volumétrie.",
            "Exécuter, avec l'équipe applicative, les comptages de contrôle "
            "sur les tables de référence et les comparer à DC1.",
        ],
        controls=[
            "Service démarré, journal de démarrage exempt d'erreur.",
            "Instance accessible sur l'adresse et le port attendus de DC2.",
            "Toutes les bases de données attendues sont présentes ; la "
            "volumétrie est cohérente avec la source.",
            "Comptages de contrôle validés par l'équipe applicative.",
            "Aucun mécanisme de DC2 n'écrit vers une destination de DC1.",
        ],
        exits=[
            "Base de données DC2 en service et déclarée conforme.",
            "Résultats des contrôles d'intégrité consignés.",
            "Surveillance de l'espace disque en place, notamment pour les "
            "journaux de transactions.",
        ],
    )

    phase_block(
        document,
        "5.3", "Phase 3 — Déploiement et validation des applications sur DC2",
        "Équipe Applicative, avec appui de l'équipe Base de données",
        "Démontrer que la chaîne applicative fonctionne sur le site DC2, en "
        "s'appuyant sur la base de données restaurée.",
        actions=[
            "Déployer les composants applicatifs sur DC2 conformément à la "
            "procédure de déploiement en vigueur.",
            "Configurer les chaînes de connexion des applications DC2 vers la "
            "base de données locale de DC2.",
            "Autoriser les réseaux applicatifs de DC2 dans la configuration "
            "d'accès de la base de données, puis contrôler la validité du "
            "fichier d'autorisations.",
            "Maintenir l'étanchéité entre les deux sites : aucune application "
            "de DC2 ne se connecte à la base de données de DC1 pendant cette "
            "phase.",
            "Exécuter le jeu de recette défini par le client et consigner les "
            "résultats.",
        ],
        controls=[
            "Composants applicatifs démarrés et fonctionnels sur DC2.",
            "Connexions applicatives visibles sur la base de données DC2 "
            "uniquement.",
            "Jeu de recette exécuté intégralement, écarts documentés.",
        ],
        exits=[
            "Recette applicative DC2 prononcée par le client.",
            "Aucune application de DC2 raccordée à la base de données de DC1.",
            "Résultats de recette annexés au journal d'exécution.",
        ],
    )

    note_box(
        document,
        "Contrainte à respecter dès la phase 3",
        [
            "La phase 4 reconstruit la base de données de DC2 à partir de "
            "celle de DC1. Les écritures réalisées sur DC2 pendant les phases "
            "2 et 3 ne sont donc pas conservées.",
            "En conséquence, les scénarios de recette de la phase 3 doivent "
            "être rejouables à l'identique, et aucune donnée de référence ne "
            "doit exister uniquement sur DC2.",
        ],
    )


def section_jalon(document):
    document.add_heading("5.4 Point de décision avant la phase 4", level=2)
    para(document,
         "Ce jalon est formel : il est validé par écrit par le client, "
         "l'équipe base de données et l'équipe applicative avant tout "
         "engagement de la phase 4.")

    document.add_heading("Élément déterminant", level=3)
    para(document,
         "La réplication de flux est un mécanisme physique : le site "
         "secondaire est une copie binaire du site principal, démarrée en "
         "lecture seule et alimentée en continu par les journaux de "
         "transactions. Une instance qui a déjà validé des transactions en "
         "lecture et écriture, ce qui est le cas de DC2 à l'issue des phases 2 "
         "et 3, ne peut pas être raccordée en l'état.")
    para(document,
         "La phase 4 comporte donc une étape de reconstruction de la base de "
         "données de DC2 à partir de DC1. Cette étape est destructive pour "
         "DC2 uniquement ; l'environnement DC1 n'est pas altéré.")

    document.add_heading("Conditions de franchissement", level=3)
    table(
        document,
        ["N°", "Condition", "Responsable"],
        [
            ["1", "Recette applicative de la phase 3 prononcée et documentée",
             "Client / Applicatif"],
            ["2", "Données à conserver exportées de DC2, ou recréées sur DC1",
             "Applicatif"],
            ["3", "Applications DC2 arrêtées, aucune session cliente "
                  "résiduelle sur la base DC2", "Applicatif"],
            ["4", "Sauvegarde récente de DC1 disponible et vérifiée",
             "Infrastructure"],
            ["5", "Méthode de reconstruction de DC2 arrêtée",
             "Base de données"],
            ["6", "Fenêtre d'intervention de la phase 4 confirmée",
             "Client"],
            ["7", "Plan de retour arrière communiqué et compris",
             "Base de données"],
        ],
        widths=[1.2, 10.9, 4.5],
    )


def section_phases_fin(document):
    phase_block(
        document,
        "5.5", "Phase 4 — Activation et validation de la réplication DC1 "
               "vers DC2",
        "Équipe Base de données, sur les deux sites",
        "Faire de la base de données DC2 un réplica en lecture seule, "
        "alimenté en continu par DC1, et démontrer que le flux est "
        "opérationnel.",
        actions=[
            "Préparer le site principal : paramètres de journalisation "
            "adaptés à la réplication, compte de réplication dédié, "
            "autorisation d'accès depuis l'adresse de DC2, mécanisme de "
            "rétention des journaux garantissant que DC1 ne recycle aucun "
            "journal non encore transmis.",
            "Confirmer l'arrêt des applications DC2 et l'absence de session "
            "cliente sur la base de données DC2.",
            "Arrêter proprement l'instance DC2 ; un arrêt incomplet compromet "
            "l'étape suivante.",
            "Reconstituer la base de données DC2 en copie physique de DC1, "
            "selon la méthode retenue au point de décision. L'emplacement de "
            "données existant est conservé sous forme de sauvegarde locale "
            "jusqu'à la validation de la phase.",
            "Déclarer le mode réplica sur DC2 et renseigner le raccordement "
            "au site principal, en identifiant explicitement le réplica pour "
            "faciliter sa supervision.",
            "Démarrer l'instance DC2 et contrôler son journal de démarrage.",
            "Arrêter les applications DC2, ou les raccorder à la base de "
            "données de DC1 si leur maintien en service est souhaité. Aucune "
            "application ne doit tenter d'écrire sur DC2.",
        ],
        controls=[
            "Instance DC2 en mode restauration : une tentative d'écriture est "
            "refusée. Ce refus constitue la preuve du fonctionnement en "
            "lecture seule.",
            "Flux de réplication actif entre DC1 et DC2, réplica visible "
            "depuis le site principal.",
            "Retard de réplication faible et stable.",
            "Test de propagation concluant : une donnée créée sur DC1 dans une "
            "base technique est visible sur DC2, puis supprimée.",
            "Historique de transactions identique sur les deux sites ; une "
            "divergence signifie que DC2 a démarré en mode autonome et impose "
            "de reprendre la reconstruction.",
            "Mécanisme de rétention des journaux actif et consommé par DC2.",
        ],
        exits=[
            "Réplication DC1 vers DC2 active et vérifiée sur les deux sites.",
            "Aucune application raccordée en écriture à DC2.",
            "Supervision en place : retard de réplication, présence du "
            "réplica, rétention des journaux, espace disque.",
            "Sauvegarde locale de l'ancien emplacement de données conservée "
            "jusqu'à la fin de la phase 5.",
        ],
    )

    phase_block(
        document,
        "5.6", "Phase 5 — Tests de failover et de failback",
        "Équipe Base de données et équipe Applicative",
        "Démontrer que l'environnement UAT peut être basculé sur DC2, puis "
        "ramené sur DC1, et mesurer les durées associées.",
        actions=[
            "Vérifier les conditions d'entrée : retard de réplication proche "
            "de zéro, sauvegarde récente de DC1, plan de bascule des accès "
            "applicatifs prêt, équipes mobilisées.",
            "Geler l'activité applicative sur DC1 pour la durée du test.",
            "Exécuter la bascule de la base de données selon le scénario "
            "retenu, puis contrôler le rôle effectif de chaque site.",
            "Basculer les accès applicatifs vers le site devenu actif. En "
            "l'absence d'adresse virtuelle ou de répartiteur de connexions, "
            "cette opération est manuelle et conditionne la réussite du test.",
            "Exécuter les tests métier sur le site actif.",
            "Réaliser le retour vers DC1, puis basculer à nouveau les accès "
            "applicatifs.",
            "Consigner les durées d'indisponibilité observées à l'aller et au "
            "retour.",
        ],
        controls=[
            "Un seul site en écriture à tout instant. En cas de bascule "
            "consécutive à un incident, le site défaillant est isolé avant "
            "toute promotion du site secondaire.",
            "Tests métier concluants sur le site actif.",
            "Après retour : un site principal, un réplica, historique de "
            "transactions identique, retard nul.",
            "Aucune erreur d'écriture applicative liée à un raccordement "
            "resté sur le réplica.",
        ],
        exits=[
            "Failover réalisé et validé, failback réalisé et validé.",
            "Bascule des accès applicatifs effectuée dans les deux sens.",
            "Durées d'indisponibilité mesurées et consignées.",
            "Topologie cible rétablie : DC1 principal, DC2 réplica.",
        ],
    )

    document.add_heading("Distinction des scénarios de bascule", level=3)
    para(document,
         "Ces trois opérations répondent à des situations différentes et se "
         "testent séparément.")
    table(
        document,
        ["Scénario", "Situation simulée", "Condition d'exécution"],
        [
            ["Bascule planifiée",
             "Opération de maintenance programmée",
             "Les deux instances sont en fonctionnement. Le site principal "
             "n'est pas arrêté par avance : il est arrêté au moment opportun, "
             "après contrôle du rattrapage des journaux."],
            ["Bascule sur incident",
             "Perte du site principal",
             "Le site principal est arrêté et isolé avant toute promotion du "
             "site secondaire, afin d'exclure une double activité en "
             "écriture."],
            ["Retour (failback)",
             "Rétablissement de la situation nominale",
             "Le site initial est réintégré comme réplica, par "
             "resynchronisation ou par nouvelle copie physique selon "
             "l'ampleur de la divergence."],
        ],
        widths=[3.2, 4.4, 9.0],
    )

    phase_block(
        document,
        "5.7", "Phase 6 — Passage des applications DC2 en mode inactif",
        "Équipe Applicative, avec contrôle de l'équipe Base de données",
        "Placer le site DC2 en état de repos, prêt à reprendre le service, "
        "sans consommation de ressources ni risque d'écriture involontaire.",
        actions=[
            "Arrêter les composants applicatifs de DC2 et désactiver leur "
            "démarrage automatique. Un service redémarrant seul après un "
            "redémarrage de machine constitue le principal risque d'écriture "
            "involontaire.",
            "Neutraliser les tâches planifiées applicatives de DC2 : "
            "traitements par lots, purges, exports, connecteurs.",
            "Contrôler l'absence de session cliente sur la base de données "
            "DC2.",
            "Maintenir l'instance de base de données DC2 démarrée et en "
            "réplication.",
            "Supprimer la sauvegarde locale de l'ancien emplacement de données "
            "conservée en phase 4, et définir une politique de rétention pour "
            "les journaux archivés de DC2.",
        ],
        controls=[
            "Composants applicatifs DC2 arrêtés, démarrage automatique "
            "désactivé.",
            "Aucune session cliente sur la base de données DC2.",
            "Instance DC2 démarrée, en lecture seule, réplication active.",
            "Applications de l'UAT raccordées à DC1 uniquement.",
        ],
        exits=[
            "Site DC2 en attente, réplication maintenue et supervisée.",
            "Supervision active sur le retard de réplication, la présence du "
            "réplica, la rétention des journaux et l'espace disque.",
            "Journal d'exécution complété et transmis au client.",
        ],
    )

    note_box(
        document,
        "Précision sur l'état inactif",
        [
            "Le mode inactif porte sur les applications, non sur la base de "
            "données. L'instance DC2 reste démarrée : son arrêt "
            "interromprait la consommation des journaux de transactions et "
            "provoquerait, sur DC1, une rétention croissante pouvant saturer "
            "l'espace disque.",
        ],
    )


def section_rollback(document):
    document.add_heading("6. Plan de retour arrière", level=1)
    para(document,
         "Chaque phase dispose d'un point de retour identifié. "
         "L'environnement DC1 n'est modifié qu'en phase 4, par l'ajout des "
         "autorisations d'accès et du mécanisme de rétention destinés à DC2, "
         "et en phase 5 par la bascule de rôle. Ces modifications sont "
         "réversibles.")
    table(
        document,
        ["Phase", "Point de retour", "Action de retour arrière"],
        [
            ["1", "Avant restauration",
             "Supprimer les machines restaurées. DC1 n'est pas modifié."],
            ["2", "Instance DC2 non démarrée",
             "Restaurer à nouveau depuis le même point de restauration."],
            ["3", "Base DC2 en service",
             "Redéployer les composants applicatifs ; la base reste "
             "utilisable en l'état."],
            ["4", "Avant démarrage du réplica",
             "Rétablir la sauvegarde locale de l'emplacement de données et "
             "retirer la déclaration de mode réplica : DC2 redevient une "
             "instance autonome."],
            ["4", "Après échec de la méthode de reconstruction retenue",
             "Basculer sur une copie physique complète depuis DC1."],
            ["5", "Après bascule planifiée",
             "Exécuter la bascule inverse depuis le site devenu réplica."],
            ["5", "Après bascule sur incident",
             "Réintégrer le site initial comme réplica, par resynchronisation "
             "ou par nouvelle copie physique."],
            ["6", "Après arrêt applicatif",
             "Redémarrer les composants en les raccordant au site actif."],
        ],
        widths=[1.4, 4.8, 10.4],
    )


def section_risques(document):
    document.add_heading("7. Risques et mesures de maîtrise", level=1)
    table(
        document,
        ["Risque", "Conséquence", "Mesure de maîtrise", "Phase"],
        [
            ["Conflit d'adresses entre les machines restaurées et celles de "
             "DC1",
             "Connexions applicatives imprévisibles, écritures sur le mauvais "
             "site",
             "Restauration en réseau isolé, remap IP avant tout raccordement "
             "au réseau de production", "1"],
            ["Démarrage automatique de la base restaurée avant contrôle",
             "Écritures vers des destinations partagées avec DC1, historique "
             "de transactions divergent",
             "Démarrage automatique désactivé dès la restauration ; premier "
             "démarrage encadré", "1, 2"],
            ["Mécanismes hérités de DC1 pointant vers des destinations "
             "partagées",
             "Altération des archives ou des sauvegardes de l'environnement "
             "source",
             "Revue et neutralisation avant le premier démarrage", "2"],
            ["Données de référence créées uniquement sur DC2",
             "Perte de ces données lors de la reconstruction",
             "Scénarios de recette rejouables ; export ou recréation sur DC1 "
             "avant le point de décision", "3, 4"],
            ["Application écrivant sur le réplica",
             "Échec des traitements, ou perte de données en cas d'écriture sur "
             "un site non répliqué",
             "Arrêt ou repointage des applications DC2, désactivation du "
             "démarrage automatique, contrôle des sessions", "4, 6"],
            ["Rétention non maîtrisée des journaux de transactions",
             "Saturation de l'espace disque et arrêt de l'instance",
             "Supervision de la rétention et de l'espace disque dès "
             "l'activation de la réplication", "4, 5, 6"],
            ["Double activité en écriture lors d'une bascule sur incident",
             "Écritures divergentes non réconciliables",
             "Isolement du site défaillant avant toute promotion du site "
             "secondaire", "5"],
            ["Bascule des accès applicatifs incomplète",
             "Indisponibilité applicative alors que la base est disponible",
             "Inventaire exhaustif des chaînes de connexion établi en "
             "prérequis, bascule tracée point par point", "5"],
            ["Réplication asynchrone",
             "Perte des transactions non encore transmises en cas de panne "
             "brutale du site principal",
             "Contrôle du retard avant toute bascule planifiée ; niveau de "
             "perte acceptable arbitré avec le client", "4, 5"],
        ],
        widths=[4.0, 4.0, 6.4, 2.2],
    )


def section_annexes(document):
    document.add_heading("8. Annexes", level=1)

    document.add_heading("8.1 Paramètres à compléter avant exécution",
                         level=2)
    para(document,
         "Ces valeurs sont renseignées et validées lors de la réunion de "
         "lancement, puis annexées à la présente procédure.")
    table(
        document,
        ["Paramètre", "Valeur DC1", "Valeur DC2"],
        [
            ["Adresse de la base de données", "", ""],
            ["Port d'écoute", "", ""],
            ["Nom du service de base de données", "", ""],
            ["Emplacement des données", "", ""],
            ["Compte d'administration de la base", "", ""],
            ["Compte dédié à la réplication", "", ""],
            ["Destination d'archivage des journaux", "", ""],
            ["Réseaux applicatifs autorisés", "", ""],
        ],
        widths=[7.0, 4.8, 4.8],
    )

    document.add_heading("8.2 Journal d'exécution", level=2)
    table(
        document,
        ["Phase", "Date et heure", "Intervenant", "Résultat", "Observations"],
        [
            ["1. Restauration des VM et remap IP", "", "", "", ""],
            ["2. Mise en service de la base DC2", "", "", "", ""],
            ["3. Déploiement et recette applicative", "", "", "", ""],
            ["Point de décision", "", "", "", ""],
            ["4. Réplication DC1 vers DC2", "", "", "", ""],
            ["5. Failover", "", "", "", ""],
            ["5. Failback", "", "", "", ""],
            ["6. Applications inactives", "", "", "", ""],
        ],
        widths=[5.0, 2.8, 2.8, 2.2, 3.8],
    )

    document.add_heading("8.3 Glossaire", level=2)
    table(
        document,
        ["Terme", "Définition"],
        [
            ["Réplication de flux",
             "Transmission continue des journaux de transactions du site "
             "principal vers le site secondaire, qui les rejoue pour rester "
             "une copie à jour."],
            ["Réplica",
             "Instance de base de données alimentée par réplication, "
             "accessible en lecture seule."],
            ["Retard de réplication",
             "Écart entre l'état du site principal et celui du réplica, "
             "exprimé en volume de journaux ou en durée."],
            ["Bascule planifiée",
             "Inversion volontaire des rôles entre les deux sites, les deux "
             "instances étant en fonctionnement."],
            ["Failover",
             "Promotion du site secondaire en site principal, consécutive à "
             "la perte du site principal."],
            ["Failback",
             "Retour à la répartition des rôles initiale après un failover ou "
             "une bascule planifiée."],
            ["Isolement",
             "Mise hors service et hors réseau d'un site, afin d'exclure "
             "toute écriture concurrente pendant une bascule."],
            ["Remap IP",
             "Réaffectation du plan d'adressage réseau d'une machine "
             "restaurée, pour l'intégrer au site cible."],
        ],
        widths=[4.0, 12.6],
    )

    document.add_heading("8.4 Validation du document", level=2)
    table(
        document,
        ["Rôle", "Nom", "Date", "Signature"],
        [
            ["Client — représentant du projet", "", "", ""],
            ["Responsable Infrastructure", "", "", ""],
            ["Responsable Base de données", "", "", ""],
            ["Responsable Applicatif", "", "", ""],
        ],
        widths=[6.0, 4.0, 2.6, 4.0],
    )


def main():
    document = Document()
    setup_styles(document)
    section = setup_page(document)
    build_header(section, TITLE)
    build_footer(section, REFERENCE)

    document.core_properties.title = TITLE
    document.core_properties.subject = SUBTITLE
    document.core_properties.category = "Procédure technique"
    document.core_properties.comments = REFERENCE

    cover(document)
    section_objet(document)
    section_sequence(document)
    section_roles(document)
    section_prerequis(document)
    section_phases(document)
    section_jalon(document)
    section_phases_fin(document)
    section_rollback(document)
    section_risques(document)
    section_annexes(document)

    document.save(OUTPUT)
    print(f"Document genere : {OUTPUT}")


if __name__ == "__main__":
    main()
